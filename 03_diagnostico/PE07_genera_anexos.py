"""PE07 · Genera los anexos A-F y las figuras del informe — DICOSUR (Taller 07, SI-886).

  A  anexo_A_pestel.xlsx                 PESTEL con evidencia (Paso A)
  B  anexo_B_matrices_efi_efe.xlsx       EFI y EFE con fórmulas de ponderación (Paso B)
  C  anexo_C_foda_cruzado.png            FODA cruzado en cuadrantes (Paso C)
  D  anexo_D_estrategias_priorizadas.xlsx  priorización (Paso C)
  E  anexo_E_trazabilidad.xlsx           matriz de trazabilidad (Paso D)
  F  anexo_F_secciones_3_3_a_3_5.pdf     secciones 3.3 a 3.5 del PETI (Paso D)

Requiere haber ejecutado antes PE02_matrices.py y PE04_prioriza_estrategias.py.
El PDF del anexo F se exporta con Microsoft Word (Windows); sin Word queda el .docx.

Uso (desde la raíz del repositorio):
    python 03_diagnostico/PE07_genera_anexos.py
"""
import re
import subprocess
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
S07 = BASE.parent / "docs" / "evidencias" / "S07"
SALIDAS = S07 / "salidas"
SALIDAS.mkdir(parents=True, exist_ok=True)

AZUL = "1F2D5C"
CAB = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor=AZUL)
BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
AJUSTE = Alignment(wrap_text=True, vertical="top")

pestel = pd.read_csv(BASE / "PE01_pestel.csv")
efi = pd.read_csv(BASE / "PE02_matriz_efi.csv")
efe = pd.read_csv(BASE / "PE02_matriz_efe.csv")
foda = pd.read_csv(BASE / "PE03_foda_cruzado.csv")
prior = pd.read_csv(BASE / "PE04_estrategias_priorizadas.csv")
traz = pd.read_csv(BASE / "PE05_trazabilidad.csv")


# ----------------------------------------------------------------- Excel
def hoja(ws, df, anchos, titulo):
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13, color=AZUL)
    ws.append(["DICOSUR · PETI v0.7 · Confidencial · Taller 07 SI-886"])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(list(df.columns))
    for c in ws[4]:
        c.font, c.fill, c.border, c.alignment = CAB, FONDO, BORDE, AJUSTE
    for fila in df.itertuples(index=False):
        ws.append(list(fila))
        for c in ws[ws.max_row]:
            c.border, c.alignment = BORDE, AJUSTE
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def guardar(wb, nombre):
    wb.save(S07 / nombre)
    print(f"  docs/evidencias/S07/{nombre}")


def anexo_a():
    wb = Workbook()
    hoja(wb.active, pestel, [8, 14, 28, 45, 32, 45, 8, 10, 14, 40, 9], "Anexo A · Análisis PESTEL con evidencia")
    wb.active.title = "PESTEL"
    guardar(wb, "anexo_A_pestel.xlsx")


def anexo_b():
    wb = Workbook()
    for k, (nombre, df) in enumerate((("EFI", efi), ("EFE", efe))):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.title = nombre
        cols = ["id", "tipo", "factor", "evidencia", "origen", "peso", "calificacion", "ponderado",
                "justificacion_peso", "justificacion_calificacion"]
        d = df.copy()
        d["ponderado"] = None
        hoja(ws, d[cols], [6, 12, 34, 40, 26, 8, 11, 11, 42, 42], f"Anexo B · Matriz {nombre}")
        primera, ultima = 5, 4 + len(d)
        for r in range(primera, ultima + 1):
            ws[f"H{r}"] = f"=F{r}*G{r}"
            ws[f"H{r}"].number_format = "0.00"
        t = ultima + 1
        ws[f"E{t}"] = "TOTAL"
        ws[f"F{t}"] = f"=SUM(F{primera}:F{ultima})"
        ws[f"H{t}"] = f"=SUM(H{primera}:H{ultima})"
        ws[f"I{t}"] = f'=IF(H{t}>=2.5,"Igual o sobre el promedio de 2.5","Bajo el promedio de 2.5")'
        for c in ("E", "F", "H", "I"):
            ws[f"{c}{t}"].font = Font(bold=True)
            ws[f"{c}{t}"].border = BORDE
        ws[f"H{t}"].number_format = "0.00"
        escala = ("1 debilidad mayor · 2 debilidad menor · 3 fortaleza menor · 4 fortaleza mayor" if nombre == "EFI" else
                  "Calificación = qué tan bien RESPONDE la organización: 1 deficiente · 2 bajo el promedio · 3 sobre el promedio · 4 superior")
        ws[f"A{t + 2}"] = "Escala: " + escala
        ws[f"A{t + 2}"].font = Font(italic=True)
    ws = wb.create_sheet("Resumen")
    ws.append(["Matriz", "Total ponderado", "Interpretación"])
    for c in ws[1]:
        c.font, c.fill, c.border = CAB, FONDO, BORDE
    n_efi, n_efe = 5 + len(efi), 5 + len(efe)
    ws.append(["EFI", f"=EFI!H{n_efi}", f"=EFI!I{n_efi}"])
    ws.append(["EFE", f"=EFE!H{n_efe}", f"=EFE!I{n_efe}"])
    ws.append(["Posición IE", "Celda VIII", "Cosechar o desinvertir: posición defensiva (ver PE02_matrices.py)"])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.border = BORDE
    ws["B2"].number_format = ws["B3"].number_format = "0.00"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    guardar(wb, "anexo_B_matrices_efi_efe.xlsx")


def anexo_d():
    wb = Workbook()
    cols = ["orden", "id", "Tipo", "Factores cruzados", "peso_factores", "puntaje", "Prioridad preliminar",
            "Proyecto candidato", "Objetivo al que apunta", "Estrategia"]
    hoja(wb.active, prior[cols], [7, 7, 6, 16, 12, 9, 12, 40, 34, 70], "Anexo D · Estrategias priorizadas")
    wb.active.title = "Priorización"
    ws = wb.create_sheet("Criterio")
    for fila in (["Tipo", "Bono", "Razón"], ["DA", 1.35, "Supervivencia: reduce debilidades ante amenazas"],
                 ["FA", 1.20, "Defensa: usa fortalezas ante amenazas"], ["DO", 1.10, "Mejora: supera debilidades"],
                 ["FO", 1.00, "Crecimiento: aprovecha oportunidades"],
                 ["Puntaje", "", "Σ pesos EFI/EFE de los factores cruzados × bono por tipo"]):
        ws.append(fila)
    for c in ws[1]:
        c.font, c.fill = CAB, FONDO
    ws.column_dimensions["C"].width = 60
    guardar(wb, "anexo_D_estrategias_priorizadas.xlsx")


def anexo_e():
    wb = Workbook()
    hoja(wb.active, traz, [48, 30, 12, 12, 38, 48], "Anexo E · Matriz de trazabilidad evidencia → factor → estrategia → objetivo → proyecto")
    wb.active.title = "Trazabilidad"
    guardar(wb, "anexo_E_trazabilidad.xlsx")


# ----------------------------------------------------------------- Figuras
def figura_pestel():
    d = pestel.copy()
    d["signo"] = d["Intensidad (1–5)"] * d["Tipo (O/A)"].map({"O": 1, "A": -1})
    d = d.iloc[::-1]
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.barh([f"{i} · {dim} · {f[:38]}" for i, dim, f in zip(d.id, d["Dimensión"], d.Factor)], d.signo,
            color=["#27AE60" if s > 0 else "#C0392B" for s in d.signo])
    ax.axvline(0, color="#333")
    ax.set_xlim(-5.5, 5.5)
    ax.set_xticks(range(-5, 6))
    ax.set_xticklabels([str(abs(x)) for x in range(-5, 6)])
    ax.set_xlabel("← amenaza · intensidad (1-5) · oportunidad →")
    ax.set_title("DICOSUR · factores PESTEL por intensidad", fontweight="bold")
    ax.grid(axis="x", alpha=0.3)
    fig.tight_layout()
    fig.savefig(SALIDAS / "PE01_pestel_intensidad.png", dpi=160)
    plt.close(fig)
    print("  docs/evidencias/S07/salidas/PE01_pestel_intensidad.png")


def anexo_c():
    fig = plt.figure(figsize=(16, 10))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 10)
    ax.axis("off")
    ax.text(8, 9.65, "DICOSUR · Matriz FODA cruzado (12 estrategias, 3 por cuadrante)", ha="center", fontsize=18,
            fontweight="bold", color="#16285C")

    def caja(x, y, w, h, color, titulo, texto, tam=10.5):
        ax.add_patch(plt.Rectangle((x, y), w, h, facecolor=color, edgecolor="#16285C", linewidth=1.4))
        ax.text(x + 0.15, y + h - 0.15, titulo, ha="left", va="top", fontsize=tam + 1.5, fontweight="bold", color="#16285C")
        ax.text(x + 0.15, y + h - 0.55, texto, ha="left", va="top", fontsize=tam, linespacing=1.35)

    def lista(df, ancho):
        return "\n".join(textwrap.fill(f"{r.id}  {r.factor}", ancho, subsequent_indent="      ") for r in df.itertuples())

    def estrategias(tipo):
        filas = foda[foda.Tipo == tipo]
        return "\n".join(textwrap.fill(f"{r.id}  [{r._2}]  {r._5}", 58, subsequent_indent="        ")
                         for r in filas.itertuples(index=False))

    fort, deb = efi[efi.id.str.startswith("F")], efi[efi.id.str.startswith("D")]
    opo, ame = efe[efe.id.str.startswith("O")], efe[efe.id.str.startswith("A")]
    leyenda = "Factores internos (EFI)\ncruzados con factores\nexternos (EFE).\n\nFormato de cada estrategia:\nid [códigos] proyecto"
    caja(0.2, 6.3, 3.9, 3.0, "#E8F1FB", "PETI v0.7", leyenda)
    caja(4.3, 6.3, 5.7, 3.0, "#D5E8D4", "OPORTUNIDADES", lista(opo, 52))
    caja(10.2, 6.3, 5.6, 3.0, "#F8CECC", "AMENAZAS", lista(ame, 60), 9.5)
    caja(0.2, 3.2, 3.9, 3.0, "#DAE8FC", "FORTALEZAS", lista(fort, 36), 9.5)
    caja(0.2, 0.1, 3.9, 3.0, "#FFF2CC", "DEBILIDADES", lista(deb, 36), 8.8)
    caja(4.3, 3.2, 5.7, 3.0, "#FFFFFF", "FO · fortalezas para aprovechar oportunidades", estrategias("FO"))
    caja(10.2, 3.2, 5.6, 3.0, "#FFFFFF", "FA · fortalezas para neutralizar amenazas", estrategias("FA"))
    caja(4.3, 0.1, 5.7, 3.0, "#FFFFFF", "DO · superar debilidades con oportunidades", estrategias("DO"))
    caja(10.2, 0.1, 5.6, 3.0, "#FFFFFF", "DA · reducir debilidades ante amenazas", estrategias("DA"))
    fig.savefig(S07 / "anexo_C_foda_cruzado.png", dpi=150)
    plt.close(fig)
    print("  docs/evidencias/S07/anexo_C_foda_cruzado.png")


# ----------------------------------------------------------------- Anexo F
def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def texto_rico(par, texto):
    for i, trozo in enumerate(re.split(r"\*\*", texto)):
        for j, sub in enumerate(re.split(r"`", trozo)):
            run = par.add_run(sub)
            run.bold = i % 2 == 1
            if j % 2 == 1:
                run.font.name = "Consolas"


def md_a_docx(doc, md):
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i].rstrip()
        if l.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(l[2:])
            r.italic = True
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif l.startswith("## "):
            doc.add_heading(l[3:], level=1)
        elif l.startswith("### "):
            doc.add_heading(l[4:], level=2)
        elif l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lineas[i].strip()):
                    filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(filas), cols=len(filas[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for a, fila in enumerate(filas):
                for b, val in enumerate(fila):
                    c = t.cell(a, b)
                    c.text = ""
                    texto_rico(c.paragraphs[0], val)
                    for r in c.paragraphs[0].runs:
                        r.font.size = Pt(7.5)
                        if a == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    if a == 0:
                        sombrear(c, AZUL)
            doc.add_paragraph()
            continue
        elif l.startswith("- "):
            texto_rico(doc.add_paragraph(style="List Bullet"), l[2:])
        elif l.strip():
            texto_rico(doc.add_paragraph(), l)
        i += 1


def anexo_f():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)
    t = doc.add_paragraph()
    r = t.add_run("Anexo F · PETI DICOSUR v0.7 · Secciones 3.3 a 3.5")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x1F, 0x2D, 0x5C)
    doc.add_paragraph("Taller 07 · SI-886 Planeamiento Estratégico de TI · Universidad Privada de Tacna · Documento confidencial")
    for nombre in ("3.3_pestel.md", "3.4_foda.md", "3.5_estrategias.md"):
        md_a_docx(doc, (BASE / nombre).read_text(encoding="utf-8"))
        doc.add_page_break()
    docx = S07 / "anexo_F_secciones_3_3_a_3_5.docx"
    pdf = S07 / "anexo_F_secciones_3_3_a_3_5.pdf"
    doc.save(docx)
    ps = (f"$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
          f"$d=$w.Documents.Open('{docx}');$d.ExportAsFixedFormat('{pdf}',17);$d.Close($false);$w.Quit()")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, capture_output=True, timeout=180)
        docx.unlink()
        print("  docs/evidencias/S07/anexo_F_secciones_3_3_a_3_5.pdf")
    except Exception as exc:  # sin Word: se conserva el .docx
        print(f"  [AVISO] No se pudo exportar a PDF ({exc}); queda {docx.name}")


if __name__ == "__main__":
    print("PE07 · Anexos y figuras generados:")
    anexo_a()
    anexo_b()
    anexo_c()
    anexo_d()
    anexo_e()
    anexo_f()
    figura_pestel()
